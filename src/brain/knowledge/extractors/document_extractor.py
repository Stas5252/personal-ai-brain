"""Document extraction without sidecars or invented scanned-page contents.

Format libraries are loaded on demand. PDF OCR uses OCREngine.extract_text.
Unreadable scanned pages fail the import. Visual diagrams are not interpreted
by this text/OCR extractor; coverage is explicit in media_info.
"""
import math
import tempfile
from pathlib import Path
from typing import Optional
from src.brain.knowledge.extractors.base import BaseExtractor
from src.brain.knowledge.extractors.ocr_engine import OCREngine
from src.brain.models.file_metadata import ExtractionResult, ExtractedElement


class DocumentExtractor(BaseExtractor):
    SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md", ".html", ".htm"}

    def __init__(self, ocr_engine=None, max_pdf_pages: int = 2000,
                 max_render_pixels: int = 20_000_000, min_ocr_confidence: float = 0.5):
        if max_pdf_pages <= 0 or max_render_pixels <= 0:
            raise ValueError("PDF limits must be positive")
        if not 0 <= min_ocr_confidence <= 1:
            raise ValueError("min_ocr_confidence must be between 0 and 1")
        self._ocr = ocr_engine
        self.max_pdf_pages = max_pdf_pages
        self.max_render_pixels = max_render_pixels
        self.min_ocr_confidence = min_ocr_confidence

    def can_handle(self, extension: str, mime_type: str) -> bool:
        return extension.lower() in self.SUPPORTED_EXTS

    @staticmethod
    def _result(source_id, elements, info, page_count=None, error=None):
        elements = [e for e in elements if e.content.strip()]
        return ExtractionResult(source_id=source_id, success=bool(elements) and error is None,
                                elements=elements, raw_text="\n\n".join(e.content for e in elements),
                                page_count=page_count, media_info=info,
                                error_message=error or (None if elements else "Document contains no usable text"))

    def extract(self, file_path: Path, source_id: str = "default_source",
                derived_dir: Optional[Path] = None) -> ExtractionResult:
        p = Path(file_path)
        try:
            if not p.is_file():
                raise FileNotFoundError("Document file does not exist")
            if not self.can_handle(p.suffix, ""):
                raise ValueError("Unsupported document extension: " + p.suffix)
            ext = p.suffix.lower()
            if ext == ".pdf":
                dest = Path(derived_dir) if derived_dir else p.parent / ".derived"
                dest.mkdir(parents=True, exist_ok=True)
                return self._extract_pdf(p, source_id, dest)
            if ext == ".docx":
                return self._extract_docx(p, source_id)
            if ext == ".pptx":
                return self._extract_pptx(p, source_id)
            if ext == ".xlsx":
                return self._extract_xlsx(p, source_id)
            if ext in {".html", ".htm"}:
                return self._extract_html(p, source_id)
            return self._extract_text_md(p, source_id)
        except Exception as exc:
            return ExtractionResult(source_id=source_id, success=False,
                                    error_message="Document extraction failed: " + str(exc),
                                    media_info={"format": p.suffix.lstrip(".").upper()})

    def _extract_pdf(self, path: Path, source_id: str, derived_dir: Path) -> ExtractionResult:
        import fitz
        elements, reports, failed = [], [], []
        with fitz.open(str(path)) as document:
            if document.needs_pass:
                raise ValueError("PDF is password-protected; provide an unlocked copy")
            count = len(document)
            if count > self.max_pdf_pages:
                raise ValueError("PDF exceeds configured page limit")
            with tempfile.TemporaryDirectory(prefix="pdf-", dir=derived_dir) as scratch:
                for index in range(count):
                    number = index + 1
                    try:
                        page = document[index]
                        text = page.get_text().strip()
                        if text:
                            elements.append(ExtractedElement(element_type="paragraph", content=text,
                                                            page_number=number, heading_path=f"Страница {number}",
                                                            metadata={"provenance": "pdf_text"}))
                            reports.append({"page": number, "status": "TEXT", "visual_content_analyzed": False})
                            continue
                        if not page.get_images() and not page.get_drawings():
                            reports.append({"page": number, "status": "BLANK"})
                            continue
                        engine = self._ocr if self._ocr is not None else OCREngine.get_instance()
                        if not engine.is_available:
                            raise ValueError("OCR unavailable for scanned page")
                        dpi = 150
                        area = math.ceil(page.rect.width * dpi / 72) * math.ceil(page.rect.height * dpi / 72)
                        if area > self.max_render_pixels:
                            raise ValueError("Rendered PDF page exceeds pixel limit")
                        rendered = Path(scratch) / f"page-{number}.png"
                        page.get_pixmap(dpi=dpi).save(str(rendered))
                        text, confidence = engine.extract_text(rendered)
                        if not isinstance(text, str):
                            raise ValueError("OCR returned non-text content")
                        confidence = float(confidence)
                        if not text.strip() or not math.isfinite(confidence) or not self.min_ocr_confidence <= confidence <= 1:
                            raise ValueError("Scanned page has no reliable OCR text")
                        elements.append(ExtractedElement(element_type="ocr", content=text.strip(),
                                                        page_number=number, heading_path=f"Страница {number}",
                                                        metadata={"provenance": "ocr", "confidence": confidence,
                                                                  "needs_review": True}))
                        reports.append({"page": number, "status": "OCR", "confidence": confidence})
                    except Exception as exc:
                        failed.append(number)
                        reports.append({"page": number, "status": "FAILED", "error": str(exc)})
        info = {"format": "PDF", "pages": count, "page_reports": reports, "failed_pages": failed,
                "visual_content_analyzed": False, "content_coverage": "PARTIAL" if failed else "TEXT_AND_OCR_ONLY"}
        error = "Unreadable PDF pages: " + ", ".join(map(str, failed)) if failed else None
        return self._result(source_id, elements, info, count, error)

    @staticmethod
    def _table_text(rows):
        rows = [["" if cell is None else str(cell).strip() for cell in row] for row in rows]
        rows = [row for row in rows if any(cell != "" for cell in row)]
        if not rows:
            return ""
        headers = [cell or f"Колонка {i + 1}" for i, cell in enumerate(rows[0])]
        lines = ["Заголовки: " + " | ".join(headers)]
        for index, row in enumerate(rows[1:], 2):
            fields = [f"{headers[i] if i < len(headers) else f'Колонка {i + 1}'}: {cell}" for i, cell in enumerate(row)]
            lines.append(f"Строка {index}: " + "; ".join(fields))
        return "\n".join(lines)

    def _extract_docx(self, path: Path, source_id: str) -> ExtractionResult:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = docx.Document(str(path))
        elements, headings, tables = [], [], 0
        # Keep tables interleaved with paragraphs in original body order.
        for child in document.element.body.iterchildren():
            if child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                text = paragraph.text.strip()
                if not text:
                    continue
                style = paragraph.style.name if paragraph.style else ""
                level = int(style[8:]) if style.startswith("Heading ") and style[8:].isdigit() else None
                if level and 1 <= level <= 9:
                    headings = headings[:level - 1] + [text]
                    kind = "heading"
                else:
                    kind = "list_item" if "List" in style else "paragraph"
                elements.append(ExtractedElement(element_type=kind, content=text,
                                                 heading_path=" > ".join(headings) or None))
            elif child.tag.endswith("}tbl"):
                tables += 1
                table = Table(child, document)
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                content = self._table_text(rows)
                if content:
                    elements.append(ExtractedElement(element_type="table", content=content,
                                                     heading_path=" > ".join(headings) or None,
                                                     metadata={"table_index": tables, "rows": len(rows)}))
        return self._result(source_id, elements, {"format": "DOCX", "tables": tables,
                           "visual_content_analyzed": False, "content_coverage": "BODY_TEXT_AND_TABLES_ONLY"})

    def _extract_pptx(self, path: Path, source_id: str) -> ExtractionResult:
        import pptx
        presentation = pptx.Presentation(str(path))
        elements, empty_slides = [], []
        def parts(shapes):
            for shape in shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    yield shape.text.strip()
                if getattr(shape, "has_table", False):
                    yield self._table_text([[cell.text for cell in row.cells] for row in shape.table.rows])
                if hasattr(shape, "shapes"):
                    yield from parts(shape.shapes)
        for number, slide in enumerate(presentation.slides, 1):
            title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
            texts = [part for part in parts(slide.shapes) if part]
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            if not texts and not notes:
                empty_slides.append(number)
                continue
            if texts:
                elements.append(ExtractedElement(element_type="slide", content="\n".join(texts),
                                                 slide_number=number, heading_path=title or f"Слайд {number}",
                                                 metadata={"visual_content_analyzed": False}))
            if notes:
                elements.append(ExtractedElement(element_type="speaker_notes", content=notes,
                                                 slide_number=number, heading_path=title or f"Слайд {number}",
                                                 metadata={"visual_content_analyzed": False}))
        return self._result(source_id, elements, {"format": "PPTX", "slides": len(presentation.slides),
                           "slides_without_text": empty_slides, "visual_content_analyzed": False,
                           "content_coverage": "TEXT_TABLES_AND_NOTES_ONLY"}, len(presentation.slides))

    def _extract_xlsx(self, path: Path, source_id: str) -> ExtractionResult:
        import openpyxl
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        elements, sheets = [], list(workbook.sheetnames)
        try:
            for sheet in workbook.worksheets:
                header = None
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = ["" if v is None else str(v).strip() for v in row]
                    if not any(v != "" for v in values):
                        continue
                    if header is None:
                        header = [v or f"Колонка {i + 1}" for i, v in enumerate(values)]
                        content = "Заголовки: " + " | ".join(header)
                    else:
                        fields = [f"{header[i] if i < len(header) else f'Колонка {i + 1}'}: {v}"
                                  for i, v in enumerate(values) if v != ""]
                        content = f"Лист '{sheet.title}', строка {number}: " + "; ".join(fields)
                    elements.append(ExtractedElement(element_type="table", content=content,
                                                     sheet_name=sheet.title, heading_path=sheet.title,
                                                     metadata={"row_number": number}))
        finally:
            workbook.close()
        return self._result(source_id, elements, {"format": "XLSX", "sheets": sheets,
                           "formula_values": "CACHED_ONLY", "content_coverage": "CELL_VALUES_ONLY"})

    def _extract_html(self, path: Path, source_id: str) -> ExtractionResult:
        from bs4 import BeautifulSoup, NavigableString
        soup = BeautifulSoup(path.read_bytes(), "html.parser")
        for tag in soup(["script", "style", "nav", "header", "footer", "aside", "noscript", "svg"]):
            tag.decompose()
        title = soup.title.get_text(" ", strip=True) if soup.title else path.stem
        body = soup.body or soup
        elements, buffer = [], []
        heading = title
        def flush():
            text = " ".join(buffer).strip()
            buffer.clear()
            if text:
                elements.append(ExtractedElement(element_type="paragraph", content=text, heading_path=heading))
        def visit(node):
            nonlocal heading
            if isinstance(node, NavigableString):
                if type(node) is NavigableString and str(node).strip():
                    buffer.append(str(node).strip())
                return
            if node.name in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                flush()
                heading = node.get_text(" ", strip=True)
                if heading:
                    elements.append(ExtractedElement(element_type="heading", content=heading, heading_path=heading))
                return
            if node.name == "table":
                flush()
                content = self._table_text([[cell.get_text(" ", strip=True)
                                            for cell in row.find_all(["td", "th"], recursive=False)]
                                           for row in node.find_all("tr")])
                if content:
                    elements.append(ExtractedElement(element_type="table", content=content, heading_path=heading))
                return
            boundary = node.name in {"div", "section", "article", "p", "li", "ul", "ol", "br", "pre"}
            if boundary:
                flush()
            for child in node.children:
                visit(child)
            if boundary:
                flush()
        visit(body)
        flush()
        return self._result(source_id, elements, {"format": "HTML", "title": title,
                           "encoding": soup.original_encoding, "content_coverage": "VISIBLE_TEXT_ONLY",
                           "conversation_structure_preserved": False})

    def _extract_text_md(self, path: Path, source_id: str) -> ExtractionResult:
        text = path.read_text(encoding="utf-8-sig")
        elements, buffer, headings = [], [], []
        def flush():
            if buffer:
                elements.append(ExtractedElement(element_type="paragraph", content="\n".join(buffer),
                                                 heading_path=" > ".join(headings) or None))
                buffer.clear()
        for line in text.splitlines():
            stripped = line.strip()
            if not stripped:
                flush()
            elif path.suffix.lower() == ".md" and stripped.startswith("#") and " " in stripped:
                prefix, heading = stripped.split(" ", 1)
                if 1 <= len(prefix) <= 6 and set(prefix) == {"#"} and heading.strip():
                    flush()
                    headings = headings[:len(prefix) - 1] + [heading.strip()]
                    elements.append(ExtractedElement(element_type="heading", content=heading.strip(),
                                                     heading_path=" > ".join(headings)))
                else:
                    buffer.append(line)
            else:
                buffer.append(line)
        flush()
        return self._result(source_id, elements, {"format": path.suffix.lstrip(".").upper(),
                           "encoding": "utf-8", "content_coverage": "TEXT_ONLY"})
