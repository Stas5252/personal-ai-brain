"""Honest document/image extraction regression tests.

Run from repository root:
    python -m unittest discover -s tests/reliability -p test_honest_ingestion.py -v

Image files and Office/HTML/text parsing are real. OCR/vision are deterministic
contract doubles, NOT a claim of real Russian OCR/Gemini accuracy. PDF contract
tests use a double; additional real-PyMuPDF tests run only when installed.
No network calls, application database or user's files are used.
"""
import importlib.util
import inspect
import json
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from src.brain.knowledge.extractors.document_extractor import DocumentExtractor
from src.brain.knowledge.extractors.image_extractor import ImageExtractor
from src.brain.knowledge.extractors.ocr_engine import OCREngine
from src.brain.knowledge.extractors.vision_provider import VisionStatus, VisionAnalysisResult


class OCRDouble:
    def __init__(self, text="Пакет Портрет: 15000 рублей", confidence=0.9, available=True, error=False):
        self.text, self.confidence = text, confidence
        self.is_available, self.error, self.calls = available, error, []

    def extract_text(self, path):
        self.calls.append(Path(path))
        if self.error:
            raise RuntimeError("recognizer failed")
        if not Path(path).is_file():
            raise AssertionError("OCR must receive a real file")
        return self.text, self.confidence


class VisionDouble:
    def __init__(self, available=False, description="", status=VisionStatus.AVAILABLE, error=False):
        self.available, self.description, self.status = available, description, status
        self.error, self.calls, self.metadata = error, [], []

    def is_available(self):
        return self.available

    def analyze_image(self, path):
        self.calls.append(Path(path))
        with Image.open(path) as image:
            image.load()
            self.metadata.append({"size": image.size, "exif": dict(image.getexif()), "info": dict(image.info)})
        if self.error:
            raise RuntimeError("provider failed")
        return VisionAnalysisResult(status=self.status, description=self.description,
                                    model_name="test-only", detected_objects=["UNVERIFIED_LABEL"])


class FixtureCase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.derived = self.root / "derived"

    def image(self, name="actual.png", size=(32, 24)):
        path = self.root / name
        with Image.new("RGB", size, "white") as image:
            image.save(path)
        return path

    def document(self, name, text):
        path = self.root / name
        path.write_text(text, encoding="utf-8")
        return path


class ImageTests(FixtureCase):
    def extract(self, ocr=None, vision=None, path=None, **options):
        engine = ImageExtractor(ocr_engine=ocr if ocr is not None else OCRDouble(),
                                vision_provider=vision if vision is not None else VisionDouble(), **options)
        return engine.extract(path or self.image(), "source", self.derived)

    def test_actual_ocr_contract_and_chunk_type(self):
        self.assertTrue(callable(OCREngine.extract_text))
        ocr = OCRDouble()
        result = self.extract(ocr=ocr)
        self.assertTrue(result.success)
        self.assertEqual(result.elements[0].element_type, "ocr")
        self.assertIn("15000", result.raw_text)
        self.assertEqual(len(ocr.calls), 1)
        self.assertFalse(ocr.calls[0].exists(), "temporary image must be removed")
        self.assertEqual(result.media_info["content_coverage"], "OCR_ONLY")

    def test_sidecar_answers_are_never_used(self):
        image = self.image("hash_moodboard.png")
        for name in ["hash_moodboard.meta.json", "moodboard.meta.json", "hash_moodboard.png.vision.json"]:
            self.document(name, json.dumps({"ocr_text": "POISON", "visual_description": "POISON"}))
        self.document("hash_moodboard.png.ocr.txt", "POISON")
        result = self.extract(path=image)
        self.assertTrue(result.success)
        self.assertNotIn("POISON", result.raw_text)

    def test_metadata_only_is_failure(self):
        result = self.extract(ocr=OCRDouble(available=False))
        self.assertFalse(result.success)
        self.assertEqual(result.elements, [])
        self.assertEqual(result.raw_text, "")

    def test_unavailable_does_not_call_ocr(self):
        ocr = OCRDouble(available=False)
        self.extract(ocr=ocr)
        self.assertEqual(ocr.calls, [])

    def test_empty_ocr_is_not_success(self):
        result = self.extract(ocr=OCRDouble(text=""))
        self.assertFalse(result.success)
        self.assertEqual(result.media_info["ocr_status"], "EMPTY_OR_FAILED")

    def test_ocr_exception_is_visible(self):
        result = self.extract(ocr=OCRDouble(error=True))
        self.assertFalse(result.success)
        self.assertEqual(result.media_info["ocr_status"], "ERROR")

    def test_low_confidence_not_indexed(self):
        result = self.extract(ocr=OCRDouble(confidence=0.1))
        self.assertFalse(result.success)
        self.assertNotIn("15000", result.raw_text)
        self.assertEqual(result.media_info["ocr_status"], "LOW_CONFIDENCE")

    def test_invalid_confidence_rejected(self):
        for confidence in [float("nan"), float("inf"), -1, 2]:
            with self.subTest(confidence=confidence):
                self.assertFalse(self.extract(ocr=OCRDouble(confidence=confidence)).success)

    def test_nontext_ocr_rejected(self):
        self.assertFalse(self.extract(ocr=OCRDouble(text={"bad": "contract"})).success)

    def test_vision_only_is_explicit_interpretation(self):
        result = self.extract(ocr=OCRDouble(available=False),
                              vision=VisionDouble(True, "Белый фон и боковой свет"))
        self.assertTrue(result.success)
        self.assertEqual(result.media_info["content_coverage"], "VISION_ONLY")
        self.assertEqual(result.elements[0].element_type, "vision_description")
        self.assertEqual(result.elements[0].metadata["provenance"], "model_interpretation")
        self.assertEqual(result.media_info["detected_objects"], [])

    def test_combined_coverage(self):
        result = self.extract(vision=VisionDouble(True, "Текст на светлом фоне"))
        self.assertEqual(result.media_info["content_coverage"], "OCR_AND_VISION")
        self.assertEqual([e.element_type for e in result.elements], ["ocr", "vision_description"])

    def test_vision_error_preserves_valid_ocr_with_warning(self):
        result = self.extract(vision=VisionDouble(True, error=True))
        self.assertTrue(result.success)
        self.assertEqual(result.media_info["vision_status"], "ERROR")
        self.assertIn("Visual content was not analyzed", result.media_info["warnings"])

    def test_vision_empty_or_failed_does_not_supply_content(self):
        for provider in [VisionDouble(True, ""), VisionDouble(True, "SHOULD_NOT_APPEAR", VisionStatus.ERROR)]:
            with self.subTest(provider=provider):
                result = self.extract(ocr=OCRDouble(available=False), vision=provider)
                self.assertFalse(result.success)
                self.assertEqual(result.raw_text, "")

    def test_corrupt_image_does_not_reach_provider(self):
        path = self.document("broken.png", "not an image")
        ocr, vision = OCRDouble(), VisionDouble(True, "unused")
        self.assertFalse(self.extract(path=path, ocr=ocr, vision=vision).success)
        self.assertEqual(ocr.calls + vision.calls, [])

    def test_pixel_limit(self):
        self.assertFalse(self.extract(max_pixels=10).success)

    def test_multiframe_tiff_is_not_silently_truncated(self):
        path = self.root / "pages.tiff"
        with Image.new("RGB", (12, 12), "white") as first, Image.new("RGB", (12, 12), "red") as second:
            first.save(path, save_all=True, append_images=[second])
        result = self.extract(path=path)
        self.assertFalse(result.success)
        self.assertIn("Multi-frame", result.error_message)

    def test_orientation_and_exif_privacy(self):
        path = self.root / "portrait.jpg"
        with Image.new("RGB", (20, 10), "white") as image:
            exif = Image.Exif()
            exif[274] = 6
            exif[315] = "PRIVATE_AUTHOR"
            image.save(path, exif=exif)
        vision = VisionDouble(True, "Проверка")
        result = self.extract(path=path, vision=vision)
        self.assertTrue(result.success)
        self.assertEqual(vision.metadata[0]["size"], (10, 20))
        self.assertEqual(vision.metadata[0]["exif"], {})
        self.assertNotIn("exif", vision.metadata[0]["info"])
        self.assertFalse(vision.calls[0].exists())

    def test_source_id_is_not_a_path(self):
        result = ImageExtractor(OCRDouble(), VisionDouble()).extract(
            self.image(), "../../not-a-path", self.derived)
        self.assertTrue(result.success)
        self.assertEqual(list(self.derived.iterdir()), [])

    def test_unsupported_extension(self):
        self.assertFalse(self.extract(path=self.document("file.json", "{}" )).success)


class DocumentTests(FixtureCase):
    def test_utf8_bom_and_multiline_chat(self):
        path = self.document("chat.txt", "\ufeffАня: Сколько стоит?\nЯ: 15000 рублей")
        result = DocumentExtractor().extract(path)
        self.assertTrue(result.success)
        self.assertEqual(result.raw_text, "Аня: Сколько стоит?\nЯ: 15000 рублей")

    def test_empty_text_not_success(self):
        self.assertFalse(DocumentExtractor().extract(self.document("empty.txt", " \n")).success)

    def test_invalid_encoding_not_silently_corrupted(self):
        path = self.root / "legacy.txt"
        path.write_bytes("Прайс".encode("cp1251"))
        self.assertFalse(DocumentExtractor().extract(path).success)

    def test_markdown_heading_context(self):
        path = self.document("lesson.md", "# Свет\n\n## Окно\nМягкий свет\n\nВторой абзац")
        result = DocumentExtractor().extract(path)
        self.assertTrue(result.success)
        self.assertEqual(result.elements[-1].heading_path, "Свет > Окно")

    def test_txt_hash_is_not_markdown(self):
        result = DocumentExtractor().extract(self.document("text.txt", "#123 заказ\n# клиент"))
        self.assertEqual(result.elements[0].element_type, "paragraph")
        self.assertIn("#123", result.raw_text)

    def test_html_div_messages_and_no_duplicates(self):
        path = self.document("chat.html", '<meta charset="utf-8"><body><h1>Чат</h1><div>Аня: цена?</div><div>Я: 15000</div><ul><li><p>Уникальная фраза</p></li></ul><script>POISON</script><!-- SECRET --></body>')
        result = DocumentExtractor().extract(path)
        self.assertTrue(result.success)
        self.assertIn("Аня: цена?", result.raw_text)
        self.assertIn("Я: 15000", result.raw_text)
        self.assertEqual(result.raw_text.count("Уникальная фраза"), 1)
        self.assertNotIn("POISON", result.raw_text)
        self.assertNotIn("SECRET", result.raw_text)
        self.assertFalse(result.media_info["conversation_structure_preserved"])

    def test_docx_original_table_order(self):
        import docx
        path = self.root / "price.docx"
        doc = docx.Document()
        doc.add_heading("Прайс", 1)
        doc.add_paragraph("ДО ТАБЛИЦЫ")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Пакет", "Цена"
        table.cell(1, 0).text, table.cell(1, 1).text = "Портрет", "15000"
        doc.add_paragraph("ПОСЛЕ ТАБЛИЦЫ")
        doc.save(path)
        result = DocumentExtractor().extract(path)
        self.assertTrue(result.success)
        self.assertLess(result.raw_text.index("ДО ТАБЛИЦЫ"), result.raw_text.index("15000"))
        self.assertLess(result.raw_text.index("15000"), result.raw_text.index("ПОСЛЕ ТАБЛИЦЫ"))
        self.assertEqual(result.elements[2].heading_path, "Прайс")

    def test_single_row_table_not_lost(self):
        text = DocumentExtractor._table_text([["Портрет", 15000]])
        self.assertIn("Портрет", text)
        self.assertIn("15000", text)

    def test_xlsx_zero_false_and_physical_rows(self):
        import openpyxl
        path = self.root / "price.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Условия"
        sheet.append(["Цена", "Оплачено"])
        sheet.append([None, None])
        sheet.append([0, False])
        book.save(path)
        book.close()
        result = DocumentExtractor().extract(path)
        self.assertTrue(result.success)
        self.assertIn("Цена: 0", result.raw_text)
        self.assertIn("Оплачено: False", result.raw_text)
        self.assertEqual(result.elements[-1].metadata["row_number"], 3)
        self.assertEqual(result.elements[-1].sheet_name, "Условия")

    def test_pptx_table_and_notes_once(self):
        import pptx
        from pptx.util import Inches
        path = self.root / "lesson.pptx"
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Урок света"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(2)).table
        table.cell(0, 0).text, table.cell(0, 1).text = "Схема", "Угол"
        table.cell(1, 0).text, table.cell(1, 1).text = "Боковой", "45"
        slide.notes_slide.notes_text_frame.text = "УНИКАЛЬНАЯ ЗАМЕТКА"
        presentation.save(path)
        result = DocumentExtractor().extract(path)
        self.assertTrue(result.success)
        self.assertIn("Угол: 45", result.raw_text)
        self.assertEqual(result.raw_text.count("УНИКАЛЬНАЯ ЗАМЕТКА"), 1)
        self.assertEqual(result.elements[0].slide_number, 1)

    def test_corrupt_office_and_missing_file_fail(self):
        for name in ["bad.docx", "bad.pptx", "bad.xlsx"]:
            with self.subTest(name=name):
                self.assertFalse(DocumentExtractor().extract(self.document(name, "bad")).success)
        self.assertFalse(DocumentExtractor().extract(self.root / "missing.txt").success)

    def test_no_production_sidecar_lookup(self):
        for cls in [DocumentExtractor, ImageExtractor]:
            source = inspect.getsource(cls)
            self.assertNotIn("tests/pilot_corpus", source)
            self.assertNotIn(".meta.json", source)
            self.assertNotIn("ocr_image(", source)


class PDFPageDouble:
    def __init__(self, text="", image=True, huge=False):
        self.text, self.image = text, image
        self.rect = types.SimpleNamespace(width=100000 if huge else 100, height=100)

    def get_text(self):
        return self.text

    def get_images(self):
        return [object()] if self.image else []

    def get_drawings(self):
        return []

    def get_pixmap(self, dpi):
        def save(path):
            with Image.new("RGB", (32, 24), "white") as image:
                image.save(path)
        return types.SimpleNamespace(save=save)


class PDFDocumentDouble:
    def __init__(self, pages, locked=False):
        self.pages, self.needs_pass, self.closed = pages, locked, False

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.closed = True

    def __len__(self):
        return len(self.pages)

    def __getitem__(self, index):
        return self.pages[index]


class PDFContractTests(FixtureCase):
    def run_pdf(self, pages, ocr=None, locked=False, **options):
        document = PDFDocumentDouble(pages, locked)
        module = types.ModuleType("fitz")
        module.open = lambda path: document
        path = self.document("course.pdf", "contract fixture, not a real PDF")
        self.document("course.meta.json", json.dumps({"ocr_text": "POISON"}))
        with patch.dict("sys.modules", {"fitz": module}):
            result = DocumentExtractor(ocr_engine=ocr if ocr is not None else OCRDouble(), **options).extract(
                path, "pdf-source", self.derived)
        self.assertTrue(document.closed)
        return result

    def test_real_ocr_method_used_no_sidecar(self):
        ocr = OCRDouble()
        result = self.run_pdf([PDFPageDouble()], ocr=ocr)
        self.assertTrue(result.success)
        self.assertIn("15000", result.raw_text)
        self.assertNotIn("POISON", result.raw_text)
        self.assertEqual(result.elements[0].page_number, 1)
        self.assertEqual(result.elements[0].element_type, "ocr")
        self.assertFalse(ocr.calls[0].exists())

    def test_one_unreadable_page_fails_whole_import(self):
        result = self.run_pdf([PDFPageDouble("Native text"), PDFPageDouble()], ocr=OCRDouble(available=False))
        self.assertFalse(result.success)
        self.assertEqual(result.media_info["failed_pages"], [2])
        self.assertIn("Native text", result.raw_text)
        self.assertNotIn("Сканированная страница", result.raw_text)

    def test_low_confidence_page_fails(self):
        self.assertFalse(self.run_pdf([PDFPageDouble()], ocr=OCRDouble(confidence=0.1)).success)

    def test_blank_document_does_not_become_knowledge(self):
        result = self.run_pdf([PDFPageDouble(image=False)])
        self.assertFalse(result.success)
        self.assertEqual(result.raw_text, "")

    def test_blank_page_recorded_but_not_fabricated(self):
        result = self.run_pdf([PDFPageDouble(image=False), PDFPageDouble("Actual text")])
        self.assertTrue(result.success)
        self.assertEqual(len(result.elements), 1)
        self.assertEqual(result.elements[0].page_number, 2)
        self.assertEqual(result.media_info["page_reports"][0]["status"], "BLANK")

    def test_encrypted_document_fails(self):
        self.assertFalse(self.run_pdf([PDFPageDouble()], locked=True).success)

    def test_page_limit(self):
        self.assertFalse(self.run_pdf([PDFPageDouble(), PDFPageDouble()], max_pdf_pages=1).success)

    def test_pixel_limit_before_rendering(self):
        self.assertFalse(self.run_pdf([PDFPageDouble(huge=True)]).success)


@unittest.skipUnless(importlib.util.find_spec("fitz"), "PyMuPDF unavailable; real PDF parsing NOT verified")
class RealPDFTests(FixtureCase):
    def test_actual_text_pdf_pages(self):
        import fitz
        path = self.root / "actual.pdf"
        with fitz.open() as document:
            document.new_page().insert_text((72, 72), "Price 15000")
            document.new_page().insert_text((72, 72), "Delivery 14 days")
            document.save(path)
        result = DocumentExtractor(ocr_engine=OCRDouble(available=False)).extract(path)
        self.assertTrue(result.success)
        self.assertEqual([e.page_number for e in result.elements], [1, 2])
        self.assertIn("15000", result.raw_text)
        self.assertIn("14 days", result.raw_text)

    def test_actual_scanned_pdf_render_and_ocr_contract(self):
        import fitz
        path, image = self.root / "scan.pdf", self.image()
        with fitz.open() as document:
            page = document.new_page()
            page.insert_image(page.rect, filename=str(image))
            document.save(path)
        ocr = OCRDouble()
        result = DocumentExtractor(ocr_engine=ocr).extract(path)
        self.assertTrue(result.success)
        self.assertEqual(len(ocr.calls), 1)
        self.assertEqual(result.elements[0].element_type, "ocr")


if __name__ == "__main__":
    unittest.main()
